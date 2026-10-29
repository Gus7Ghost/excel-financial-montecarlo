import os
import openpyxl
import pytest


@pytest.fixture(scope="module")
def workbook():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    excel_path = os.path.join(base_dir, "data", "Modelo_Valoracion_Financiera_MonteCarlo.xlsx")
    assert os.path.exists(excel_path), f"El archivo Excel no existe en {excel_path}"
    wb = openpyxl.load_workbook(excel_path, data_only=False)
    return wb


@pytest.fixture(scope="module")
def workbook_values():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    excel_path = os.path.join(base_dir, "data", "Modelo_Valoracion_Financiera_MonteCarlo.xlsx")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    return wb


def test_sheet_names_structure(workbook):
    expected_sheets = [
        "01_Parametros_Supuestos",
        "02_Flujo_Caja_DCF",
        "03_Sensibilidad_Bidimensional",
        "04_Simulacion_MonteCarlo",
        "05_PowerQuery_M_Script"
    ]
    for s in expected_sheets:
        assert s in workbook.sheetnames, f"Falta la hoja requerida: {s}"


def test_dcf_formulas(workbook):
    ws_dcf = workbook["02_Flujo_Caja_DCF"]

    # Verificar que el VAN contiene la fórmula NPV
    van_formula = str(ws_dcf["C26"].value)
    assert "NPV" in van_formula or "VAN" in van_formula, "La celda C26 debe contener la fórmula de VAN/NPV"

    # Verificar que la TIR contiene la fórmula IRR
    tir_formula = str(ws_dcf["C27"].value)
    assert "IRR" in tir_formula or "TIR" in tir_formula, "La celda C27 debe contener la fórmula de TIR/IRR"

    # Verificar la decisión lógica IF
    decision_formula = str(ws_dcf["C31"].value)
    assert "IF" in decision_formula or "SI" in decision_formula, "La celda C31 debe contener la regla de decisión condicional"


def test_sensitivity_matrix_values(workbook_values):
    ws_sens = workbook_values["03_Sensibilidad_Bidimensional"]
    # Verificar dimensiones de matriz (5x5)
    for row in range(6, 11):
        for col in ["C", "D", "E", "F", "G"]:
            val = ws_sens[f"{col}{row}"].value
            assert isinstance(val, (int, float)), f"La celda {col}{row} debe contener un valor numérico"


def test_monte_carlo_distribution(workbook_values):
    ws_mc = workbook_values["04_Simulacion_MonteCarlo"]

    # Verificar estadísticos agregados
    mean_npv = ws_mc["C6"].value
    p5_var = ws_mc["C9"].value
    prob_success = ws_mc["C12"].value

    assert isinstance(mean_npv, (int, float)) and mean_npv > 0, "El VAN promedio esperado debe ser positivo"
    assert isinstance(p5_var, (int, float)), "El VaR 95% debe ser numérico"
    assert 0.0 <= prob_success <= 1.0, "La probabilidad de éxito debe estar entre 0% y 100%"

    # Verificar que existen las 5,000 iteraciones
    first_iter = ws_mc["B17"].value
    last_iter = ws_mc["B5016"].value
    assert first_iter == 1, "La primera iteración debe tener índice 1"
    assert last_iter == 5000, "La última iteración debe tener índice 5000"
